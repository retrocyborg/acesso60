import re
import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_formatted_text(paragraph, text, is_code=False):
    # Regex to handle **bold**, *italic*, and `code`
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if is_code:
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            if is_code:
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(160, 40, 40)
        else:
            run = paragraph.add_run(token)
            if is_code:
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)

def convert_md_to_docx(md_path, docx_path):
    doc = Document()

    # Set standard page margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lang = ""
    code_lines = []
    
    in_table = False
    table_rows = []

    def flush_code():
        nonlocal in_code_block, code_lang, code_lines
        if code_lines:
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Inches(6.5)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            # Add language tag if available
            if code_lang:
                tag_run = p.add_run(f"[{code_lang.upper()}]\n")
                tag_run.font.name = 'Consolas'
                tag_run.font.size = Pt(8.5)
                tag_run.bold = True
                tag_run.font.color.rgb = RGBColor(100, 116, 139)
                
            code_text = "".join(code_lines)
            run = p.add_run(code_text.rstrip())
            run.font.name = 'Consolas'
            run.font.size = Pt(9.0)
            run.font.color.rgb = RGBColor(15, 23, 42)
            
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_before = Pt(0)
            p_after.paragraph_format.space_after = Pt(4)
            
        code_lines = []
        code_lang = ""
        in_code_block = False

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Clean rows
        parsed_rows = []
        for r in table_rows:
            cells = [c.strip() for c in r.strip().strip('|').split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                continue # delimiter row
            parsed_rows.append(cells)
            
        if not parsed_rows:
            in_table = False
            table_rows = []
            return
            
        cols_count = max(len(r) for r in parsed_rows)
        tbl = doc.add_table(rows=len(parsed_rows), cols=cols_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for row_idx, r_data in enumerate(parsed_rows):
            for col_idx in range(cols_count):
                cell_text = r_data[col_idx] if col_idx < len(r_data) else ""
                cell = tbl.cell(row_idx, col_idx)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                if row_idx == 0:
                    set_cell_background(cell, "E2E8F0")
                    run = p.add_run(cell_text)
                    run.bold = True
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    add_formatted_text(p, cell_text)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)
        table_rows = []
        in_table = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                flush_table()

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal rules
        if stripped in ["---", "***", "___"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Headings
        if stripped.startswith("# ") and not stripped.startswith("## "):
            h = doc.add_heading(level=0)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        if stripped.startswith("## "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        if stripped.startswith("### "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = RGBColor(2, 132, 199)
            i += 1
            continue

        if stripped.startswith("#### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run = h.add_run(stripped[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Blockquote (> ...)
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="3B82F6"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            quote_text = re.sub(r'^>\s*', '', stripped)
            add_formatted_text(p, quote_text)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Unordered list (- or *)
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            item_text = stripped[2:].strip()
            add_formatted_text(p, item_text)
            i += 1
            continue

        # Numbered list (1. 2. etc.)
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            item_text = num_match.group(2).strip()
            add_formatted_text(p, item_text)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, stripped)
        i += 1

    if in_code_block:
        flush_code()
    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Salvo com sucesso: {docx_path}")

if __name__ == "__main__":
    md_file = r"C:\Prog\acesso60\docs\engineering\developer-reference.md"
    docx_file = r"C:\Prog\acesso60\docs\engineering\developer-reference.docx"
    convert_md_to_docx(md_file, docx_file)
